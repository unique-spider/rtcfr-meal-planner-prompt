from docx import Document
from docx.shared import Pt

doc = Document()

doc.add_heading("RTCFR Prompt - South Indian Vegetarian Weight-Loss Meal Planner", level=1)

doc.add_heading("Role", level=2)
doc.add_paragraph(
    "You are a Certified Nutritionist with 10+ years of experience designing South Indian "
    "vegetarian diets for working professionals who want sustainable weight loss."
)

doc.add_heading("Task", level=2)
doc.add_paragraph(
    "Create a 7-day South Indian vegetarian meal plan for healthy weight loss. Cover "
    "breakfast, lunch, dinner, and one snack for each day."
)

doc.add_heading("Context", level=2)
doc.add_paragraph(
    "The person is a vegetarian working professional based in Tamil Nadu who spends most of "
    "the day at an office desk with low physical activity. They want to lose weight while "
    "still eating familiar South Indian food. Meals should use simple, easily available "
    "ingredients, be office-friendly (easy to carry or eat at a desk), keep portions "
    "moderate, and avoid deep-fried items. Assume a daily target of roughly 1500-1600 "
    "calories, moderate protein, and controlled carbohydrates."
)

doc.add_heading("Few Shots", level=2)
doc.add_paragraph(
    "Follow this style for each meal entry - short, specific, with approximate portion size:"
)
p = doc.add_paragraph()
p.add_run("Example 1 - Breakfast: ").bold = True
p.add_run("2 idlis with sambar and mint chutney, approx. 250 kcal.")
p = doc.add_paragraph()
p.add_run("Example 2 - Dinner: ").bold = True
p.add_run("1 bowl vegetable kootu with 1 small millet dosa, approx. 300 kcal.")

doc.add_heading("Response Format", level=2)
doc.add_paragraph(
    "Return the plan as a clean day-wise table with these columns: Day, Breakfast, Lunch, "
    "Dinner, Snack, Approx. Calories. Use one row per day, Day 1 through Day 7. Keep each "
    "cell short and specific like the examples above."
)

doc.add_heading("Full Prompt (copy-paste ready)", level=2)
full_prompt = (
    "Act as a Certified Nutritionist with 10+ years of experience designing South Indian "
    "vegetarian diets for working professionals who want sustainable weight loss.\n\n"
    "Task: Create a 7-day South Indian vegetarian meal plan for healthy weight loss, covering "
    "breakfast, lunch, dinner, and one snack each day.\n\n"
    "Context: The person is a vegetarian working professional based in Tamil Nadu who spends "
    "most of the day at an office desk with low physical activity. They want to lose weight "
    "while still eating familiar South Indian food. Use simple, easily available ingredients, "
    "keep meals office-friendly, keep portions moderate, avoid deep-fried items, and target "
    "roughly 1500-1600 calories a day with moderate protein and controlled carbohydrates.\n\n"
    "Few Shots (match this style):\n"
    "- Breakfast: 2 idlis with sambar and mint chutney, approx. 250 kcal.\n"
    "- Dinner: 1 bowl vegetable kootu with 1 small millet dosa, approx. 300 kcal.\n\n"
    "Response Format: Return a day-wise table with columns Day, Breakfast, Lunch, Dinner, "
    "Snack, Approx. Calories, with one row per day from Day 1 to Day 7."
)
para = doc.add_paragraph()
run = para.add_run(full_prompt)
run.font.size = Pt(11)

doc.save("RTCFR_Meal_Planner_Prompt.docx")
print("saved")
